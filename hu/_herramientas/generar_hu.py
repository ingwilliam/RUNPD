"""Genera una historia de usuario en Word a partir de un JSON, usando la plantilla oficial.

Uso:
    python hu/_herramientas/generar_hu.py hu/_herramientas/<nombre>.json [--salida ruta.docx]

Por defecto el documento se guarda en hu/versiones/ con nombre versionado:
    HU_<nombre>_v<n>_<AAAAMMDD-HHMM>.docx
donde <n> es acumulado por HU (v1, v2... vn). Con --salida se respeta la ruta tal cual.

Reglas del JSON (ver ejemplo.json):
* Una sección ausente o null deja la plantilla intacta (con su texto guía).
* Los textos largos son listas de párrafos. Un párrafo que empieza con "- " se
  convierte en viñeta y **texto** se escribe en negrita.
* `capturas` (en cada historia) y `capturas_generales` incrustan imágenes PNG con
  un título; cada captura es {"archivo", "titulo", "ancho_in"} (ancho en pulgadas).
"""

import argparse
import copy
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

CARPETA_HU = Path(__file__).resolve().parent.parent
PLANTILLA = CARPETA_HU / "Anexo_Formato_Historias_Usuario.docx"
CARPETA_VERSIONES = CARPETA_HU / "versiones"

TAMANO_TEXTO = Pt(9)
CAMPOS_IDENTIFICACION = {
    "nombre_unidad": "Nombre de la unidad",
    "nombre_proyecto": "Nombre del proyecto",
    "requerimiento_aplicacion": "Requerimiento o aplicación",
    "lider_funcional": "Líder Funcional",
    "equipo_lider_funcional": "Equipo designado",
    "gerente_interno": "Gerente Interno",
    "lider_desarrollo": "Líder del equipo de desarrollo",
    "requerimiento_asociado": "Requerimiento asociado",
    "prioridad": "Prioridad",
    "modulo_hito": "Módulo/Hito",
}
FRECUENCIAS = {
    "diario": "Frecuente (Diario)",
    "ocasional": "Ocasional",
    "definida": "Con frecuencia definida",
    "a_solicitud": "A solicitud",
}
ACCESOS = {"web": "WEB", "externo": "Externo", "escritorio_interno": "Escritorio Interno"}
ANCLAS_BLOQUE_HU = (
    "Historia de usuario HU",
    "HU01 –",
    "Descripción:",
    "Precondiciones",
    "HU01_CA01",
)
ESPACIO_PARRAFO = Pt(4)
INTERLINEADO = 1.15
MARGENES_CELDA = {"top": 40, "left": 108, "bottom": 40, "right": 108}


# ---------- utilidades de bajo nivel ----------

def celdas_unicas(fila):
    """Celdas reales de una fila (python-docx repite las celdas combinadas)."""
    vistas, unicas = set(), []
    for celda in fila.cells:
        if id(celda._tc) not in vistas:
            vistas.add(id(celda._tc))
            unicas.append(celda)
    return unicas


def margenes_celda(celda):
    """Da aire interno a la celda para que el texto no quede pegado a los bordes."""
    tc_pr = celda._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for nombre, valor in MARGENES_CELDA.items():
        margen = mar.find(qn(f"w:{nombre}"))
        if margen is None:
            margen = OxmlElement(f"w:{nombre}")
            mar.append(margen)
        margen.set(qn("w:w"), str(valor))
        margen.set(qn("w:type"), "dxa")


def texto_de(elemento):
    """Texto completo de un nodo XML de la plantilla."""
    return "".join(t.text or "" for t in elemento.iter(qn("w:t")))


def primer_indice(textos, prefijo, contexto):
    for i, texto in enumerate(textos):
        if texto.startswith(prefijo):
            return i
    raise ValueError(
        f"No se encontró el ancla '{prefijo}' ({contexto}) en la plantilla. "
        f"¿Cambió el anexo de la UTDI?"
    )


def _primero(elementos, condicion, descripcion):
    for elemento in elementos:
        if condicion(elemento):
            return elemento
    raise ValueError(
        f"No se encontró en la plantilla el {descripcion}. ¿Cambió el anexo de la UTDI?"
    )


def buscar_fila(tabla, prefijo, desde=0):
    for i, fila in enumerate(tabla.rows):
        if i >= desde and fila.cells[0].text.strip().lower().startswith(prefijo.lower()):
            return i
    raise ValueError(f"No se encontró la fila '{prefijo}' en la plantilla. ¿Cambió el anexo de la UTDI?")


def es_guia(parrafo):
    """Texto guía de la plantilla: cursiva gris."""
    return any(r.italic and r.text.strip() for r in parrafo.runs)


def limpiar_runs(parrafo):
    for r in list(parrafo._p):
        if r.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:proofErr")):
            parrafo._p.remove(r)


def agregar_texto(parrafo, texto, negrita=None, tamano=TAMANO_TEXTO):
    """Agrega texto interpretando **negrita** en línea."""
    for i, trozo in enumerate(re.split(r"\*\*", texto)):
        if not trozo:
            continue
        run = parrafo.add_run(trozo)
        run.bold = True if i % 2 == 1 else negrita
        run.italic = False
        run.font.size = tamano


def poner_texto(parrafo, texto, negrita=None, tamano=TAMANO_TEXTO):
    limpiar_runs(parrafo)
    agregar_texto(parrafo, texto, negrita, tamano)


def como_lista(valor):
    if valor is None:
        return []
    return [valor] if isinstance(valor, str) else list(valor)


RAIZ = CARPETA_HU.parent


def resolver_ruta(archivo, avisos):
    """Resuelve la ruta de una captura relativa a la raíz del repo; None si no existe."""
    ruta = Path(archivo)
    if not ruta.is_absolute():
        ruta = RAIZ / ruta
    if not ruta.exists():
        avisos.append(f"no se encontró la captura {ruta}")
        return None
    return ruta


def imagen(cap, celda, avisos):
    """Crea un párrafo con la imagen de `cap` y devuelve su nodo XML (None si falta)."""
    ruta = resolver_ruta(cap.get("archivo", ""), avisos)
    if ruta is None:
        return None
    parrafo = celda.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.space_before = ESPACIO_PARRAFO
    parrafo.paragraph_format.space_after = ESPACIO_PARRAFO
    parrafo.add_run().add_picture(str(ruta), width=Inches(cap.get("ancho_in", 6)))
    return parrafo._p


def capturas(lista, celda, proto, avisos, codigo=None):
    """Devuelve los párrafos de título más la imagen de cada captura."""
    salida = []
    for cap in lista:
        prefijo = f"Captura {codigo}" if codigo else "Captura general"
        rotulo = proto.nuevo(f"{prefijo} - {cap.get('titulo', '')}", celda)
        rotulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        salida.append(rotulo._p)
        nodo = imagen(cap, celda, avisos)
        if nodo is not None:
            salida.append(nodo)
    return salida


class Prototipos:
    """Párrafos de la plantilla que se clonan para conservar su formato."""

    def __init__(self, tabla):
        objetivo = tabla.cell(buscar_fila(tabla, "Objetivo"), 0)
        self.vineta = _primero(
            objetivo.paragraphs,
            lambda p: p._p.pPr is not None and p._p.pPr.numPr is not None,
            "párrafo de viñeta (con numeración)",
        )
        self.normal = _primero(
            objetivo.paragraphs,
            lambda p: es_guia(p) and (p._p.pPr is None or p._p.pPr.numPr is None),
            "párrafo de texto guía sin numeración",
        )

    def nuevo(self, texto, contenedor):
        vineta = texto.startswith("- ")
        base = self.vineta if vineta else self.normal
        p = copy.deepcopy(base._p)
        parrafo = Paragraph(p, contenedor)
        poner_texto(parrafo, texto[2:] if vineta else texto)
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT if vineta else WD_ALIGN_PARAGRAPH.JUSTIFY
        formato = parrafo.paragraph_format
        formato.space_after = ESPACIO_PARRAFO
        formato.line_spacing = INTERLINEADO
        return parrafo


def reemplazar_contenido(celda, parrafos, proto, conservar=1):
    """Deja los primeros `conservar` párrafos (título) y reemplaza el resto por `parrafos`."""
    margenes_celda(celda)
    hijos = [h for h in celda._tc if h.tag in (qn("w:p"), qn("w:tbl"))]
    for h in hijos[conservar:]:
        celda._tc.remove(h)
    vacio = copy.deepcopy(hijos[conservar - 1])
    limpiar_runs(Paragraph(vacio, celda))
    celda._tc.append(copy.deepcopy(vacio))
    for texto in parrafos:
        celda._tc.append(proto.nuevo(texto, celda)._p)
    celda._tc.append(vacio)


# ---------- secciones ----------

def llenar_identificacion(tabla, datos, avisos):
    ident = datos.get("identificacion") or {}
    for clave, etiqueta in CAMPOS_IDENTIFICACION.items():
        valor = ident.get(clave)
        if not valor:
            avisos.append(f"identificacion.{clave} vacío")
            continue
        celda = celdas_unicas(tabla.rows[buscar_fila(tabla, etiqueta)])[1]
        poner_texto(celda.paragraphs[0], valor)
    if ident.get("fecha_solicitud"):
        celda = celdas_unicas(tabla.rows[buscar_fila(tabla, "Fecha de solicitud")])[0]
        agregar_texto(celda.paragraphs[0], " " + ident["fecha_solicitud"])


def llenar_seccion(tabla, prefijo, valor, proto):
    if valor is None:
        return
    celda = tabla.cell(buscar_fila(tabla, prefijo), 0)
    reemplazar_contenido(celda, como_lista(valor), proto)


def llenar_epica(tabla, epica, proto):
    if not epica:
        return
    titulo = f"**{epica.get('codigo', 'EP01')} – {epica.get('nombre', '')}**"
    llenar_seccion(tabla, "Casos de uso", [titulo] + como_lista(epica.get("descripcion")), proto)


def bloque_hu(celda):
    """Devuelve (hijos, inicio, fin_exclusivo) del bloque de una HU dentro de la celda."""
    hijos = [h for h in celda._tc if h.tag in (qn("w:p"), qn("w:tbl"))]
    textos = [texto_de(h) for h in hijos]
    inicio = primer_indice(textos, "HU01", "inicio del bloque de historia de usuario")
    rnf = primer_indice(textos, "Requerimientos No funcionales", "título de los requerimientos no funcionales")
    return hijos, inicio, rnf


def validar_bloque_hu(bloque):
    """Falla con un mensaje claro si el bloque de HU de la plantilla ya no tiene sus anclas."""
    textos = [texto_de(el) for el in bloque]
    for ancla in ANCLAS_BLOQUE_HU:
        if not any(texto.startswith(ancla) for texto in textos):
            raise ValueError(
                f"No se encontró el ancla '{ancla}' en el bloque de historia de usuario de la plantilla. "
                f"¿Cambió el anexo de la UTDI?"
            )


def llenar_historia(elementos, hu, celda, proto, avisos):
    """Llena una copia del bloque HU de la plantilla. `elementos` son nodos XML clonados."""
    codigo = hu["codigo"]
    for el in elementos:
        if el.tag == qn("w:tbl"):
            tabla_yo = Table(el, celda)
            for fila, clave in zip(tabla_yo.rows, ("como", "quiero", "para")):
                destino = celdas_unicas(fila)[1]
                for p in destino.paragraphs[1:]:
                    destino._tc.remove(p._p)
                poner_texto(destino.paragraphs[0], hu.get(clave) or "Por definir")
            continue
        texto = texto_de(el)
        p = Paragraph(el, celda)
        if texto.startswith("HU01 –"):
            poner_texto(p, f"{codigo} – {hu.get('nombre', '')}", tamano=Pt(9))
        elif texto.startswith("Historia de usuario HU"):
            poner_texto(p, f"Historia de usuario {codigo}", tamano=Pt(10))

    salida = []
    i = 0
    while i < len(elementos):
        el = elementos[i]
        texto = texto_de(el)
        if texto.startswith("Descripción:"):
            salida += [proto.nuevo(t, celda)._p for t in como_lista(hu.get("descripcion")) or ["Por definir"]]
        elif texto.startswith("Precondiciones"):
            precond = como_lista(hu.get("precondiciones")) or ["Ninguna."]
            salida.append(proto.nuevo("**Precondiciones:**", celda)._p)
            salida += [proto.nuevo(t, celda)._p for t in precond]
        elif texto.startswith("HU01_CA01"):
            salida += criterios(el, hu, celda, proto)
            # Omite el texto guía de criterios que sigue en la plantilla.
            while i + 1 < len(elementos):
                sig = texto_de(elementos[i + 1])
                if sig and not es_guia(Paragraph(elementos[i + 1], celda)):
                    break
                i += 1
            salida.append(vacio_de(el, celda))
        else:
            salida.append(el)
        i += 1
    salida += capturas(hu.get("capturas") or [], celda, proto, avisos, codigo=codigo)
    return salida


def vacio_de(el, celda):
    p = copy.deepcopy(el)
    limpiar_runs(Paragraph(p, celda))
    p_pr = p.find(qn("w:pPr"))
    if p_pr is not None and p_pr.find(qn("w:pStyle")) is not None:
        p_pr.remove(p_pr.find(qn("w:pStyle")))
    return p


def criterios(titulo_proto, hu, celda, proto):
    salida = []
    lista = hu.get("criterios") or []
    for n, ca in enumerate(lista, start=1):
        codigo = ca.get("codigo") or f"{hu['codigo']}_CA{n:02d}"
        titulo = copy.deepcopy(titulo_proto)
        poner_texto(Paragraph(titulo, celda), f"{codigo} – {ca.get('nombre', '')}", tamano=Pt(10))
        if n > 1:
            salida.append(vacio_de(titulo_proto, celda))
        salida.append(titulo)
        for etiqueta, clave in (("Dado", "dado"), ("Cuando", "cuando"), ("Entonces", "entonces")):
            if ca.get(clave):
                salida.append(proto.nuevo(f"**{etiqueta}** {ca[clave]}", celda)._p)
        salida += [proto.nuevo(t, celda)._p for t in como_lista(ca.get("detalle"))]
    if not lista:
        titulo = copy.deepcopy(titulo_proto)
        poner_texto(Paragraph(titulo, celda), "Por definir", tamano=Pt(10))
        salida.append(titulo)
    return salida


def llenar_historias(tabla, datos, proto, avisos):
    historias = datos.get("historias")
    if not historias:
        avisos.append("No hay historias: la sección de HU queda con el texto guía")
        return
    celda = tabla.cell(buscar_fila(tabla, "Historias de Usuario"), 0)
    margenes_celda(celda)
    hijos, inicio, rnf = bloque_hu(celda)

    # Quita el texto guía general de la sección (entre el título y HU01).
    for el in hijos[1:inicio]:
        if es_guia(Paragraph(el, celda)):
            celda._tc.remove(el)

    bloque = hijos[inicio:rnf]
    validar_bloque_hu(bloque)
    ancla = hijos[rnf]
    for n, hu in enumerate(historias):
        if n > 0:
            ancla.addprevious(vacio_de(bloque[0], celda))
        for el in llenar_historia([copy.deepcopy(e) for e in bloque], hu, celda, proto, avisos):
            ancla.addprevious(el)
    for el in bloque:
        celda._tc.remove(el)

    rnf_valor = datos.get("requerimientos_no_funcionales")
    if rnf_valor is not None:
        for el in list(ancla.itersiblings()):
            if el.tag == qn("w:p") and es_guia(Paragraph(el, celda)):
                celda._tc.remove(el)
        ultimo = ancla.getnext()  # párrafo vacío bajo el título
        for texto in como_lista(rnf_valor):
            nuevo = proto.nuevo(texto, celda)._p
            ultimo.addnext(nuevo)
            ultimo = nuevo

    for el in capturas(datos.get("capturas_generales") or [], celda, proto, avisos):
        celda._tc.append(el)


def llenar_usuarios(tabla, datos, proto):
    cantidad = datos.get("usuarios_cantidad")
    if cantidad:
        celda = celdas_unicas(tabla.rows[buscar_fila(tabla, "NÚMERO/CANTIDAD")])[1]
        for p in celda.paragraphs[1:]:
            celda._tc.remove(p._p)
        poner_texto(celda.paragraphs[0], str(cantidad))
    llenar_seccion(tabla, "IDENTIFICACIÓN DE LOS USUARIOS", datos.get("usuarios_identificacion"), proto)


def llenar_frecuencia(tabla, frecuencia, avisos):
    if not frecuencia:
        return
    if frecuencia not in FRECUENCIAS:
        avisos.append(f"frecuencia '{frecuencia}' no válida; use: {', '.join(FRECUENCIAS)}")
        return
    inicio = buscar_fila(tabla, "FRECUENCIA DE USO")
    etiqueta = tabla.cell(inicio, 0)
    for p in etiqueta.paragraphs:
        if es_guia(p):
            etiqueta._tc.remove(p._p)
    for fila in tabla.rows[inicio:]:
        celdas = celdas_unicas(fila)
        if celdas[1].text.strip().startswith(FRECUENCIAS[frecuencia]):
            p = celdas[-1].paragraphs[0]
            poner_texto(p, "X", negrita=True, tamano=Pt(10))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return


def llenar_acceso(tabla, acceso):
    if not acceso:
        return
    inicio = buscar_fila(tabla, "TIPO DE ACCESO")
    for fila in tabla.rows[inicio:inicio + len(ACCESOS)]:
        celdas = celdas_unicas(fila)
        for clave, etiqueta in ACCESOS.items():
            if celdas[1].text.strip() == etiqueta and acceso.get(clave):
                p = celdas[-1].paragraphs[0]
                poner_texto(p, acceso[clave], tamano=Pt(10))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def llenar_firmas(tabla, datos):
    for columna, clave in ((0, "solicitante"), (1, "lider_funcional_firma")):
        persona = datos.get(clave) or {}
        for etiqueta, campo in (("Nombre", "nombre"), ("Cargo", "cargo")):
            if persona.get(campo):
                fila = tabla.rows[buscar_fila(tabla, etiqueta, desde=buscar_fila(tabla, "SOLICITANTE"))]
                celda = celdas_unicas(fila)[columna]
                agregar_texto(celda.paragraphs[0], persona[campo], negrita=False, tamano=Pt(10))


def llenar_versiones(doc, version):
    if not version:
        return
    tabla = doc.tables[1]
    fila = tabla.rows[1].cells
    for celda, clave in zip(fila, ("version", "fecha", "descripcion")):
        if version.get(clave):
            poner_texto(celda.paragraphs[0], version[clave], tamano=Pt(10))


# ---------- programa ----------

def quitar_vacios_finales(doc):
    """Evita una página en blanco al final cuando el contenido crece."""
    parrafos = doc.paragraphs
    while len(parrafos) > 1 and not parrafos[-1].text.strip() and not parrafos[-2].text.strip():
        parrafos[-1]._p.getparent().remove(parrafos[-1]._p)
        parrafos = doc.paragraphs


def ruta_versionada(nombre, momento=None):
    """Devuelve HU_<nombre>_v<n>_<AAAAMMDD-HHMM>.docx con n acumulado por HU."""
    CARPETA_VERSIONES.mkdir(parents=True, exist_ok=True)
    patron = re.compile(rf"^{re.escape(nombre)}_v(\d+)_")
    existentes = [
        int(m.group(1))
        for archivo in CARPETA_VERSIONES.glob(f"{nombre}_v*")
        if (m := patron.match(archivo.stem))
    ]
    n = max(existentes, default=0) + 1
    sello = (momento or datetime.now()).strftime("%Y%m%d-%H%M")
    return CARPETA_VERSIONES / f"{nombre}_v{n}_{sello}.docx"


def generar(datos, salida):
    doc = Document(PLANTILLA)
    tabla = doc.tables[0]
    proto = Prototipos(tabla)
    avisos = []

    llenar_identificacion(tabla, datos, avisos)
    llenar_seccion(tabla, "DESCRIPCIÓN DE LA NECESIDAD", datos.get("descripcion_necesidad"), proto)
    llenar_seccion(tabla, "Objetivo", datos.get("objetivo"), proto)
    llenar_seccion(tabla, "Justificación", datos.get("justificacion"), proto)
    llenar_seccion(tabla, "Alcance", datos.get("alcance"), proto)
    llenar_epica(tabla, datos.get("epica"), proto)
    llenar_historias(tabla, datos, proto, avisos)
    llenar_usuarios(tabla, datos, proto)
    llenar_frecuencia(tabla, datos.get("frecuencia"), avisos)
    llenar_acceso(tabla, datos.get("acceso"))
    llenar_seccion(tabla, "ÁREAS INVOLUCRADAS", datos.get("areas_involucradas"), proto)
    llenar_seccion(tabla, "Observaciones adicionales", datos.get("observaciones"), proto)
    llenar_firmas(tabla, datos)
    llenar_versiones(doc, datos.get("version"))
    quitar_vacios_finales(doc)

    doc.save(salida)
    return salida, avisos


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("json", type=Path, help="Archivo JSON con el contenido de la HU")
    parser.add_argument(
        "--salida",
        type=Path,
        help="Ruta del .docx (por defecto hu/versiones/HU_<nombre>_v<n>_<AAAAMMDD-HHMM>.docx)",
    )
    args = parser.parse_args()

    if not PLANTILLA.exists():
        sys.exit(f"ERROR: no se encontró la plantilla en {PLANTILLA}")
    try:
        datos = json.loads(args.json.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        sys.exit(f"ERROR: no se pudo leer {args.json}: {error}")

    nombre = "HU_" + (datos.get("archivo") or args.json.stem)
    salida = args.salida if args.salida else ruta_versionada(nombre)
    try:
        ruta, avisos = generar(datos, salida)
    except ValueError as error:
        sys.exit(f"ERROR: {error}")
    print(f"OK: {ruta}")
    for aviso in avisos:
        print(f"AVISO: {aviso}")


if __name__ == "__main__":
    main()